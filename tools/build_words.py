"""Rebuild data/words.json and data/categories.json from the Excel word bank.

Usage:  python3 tools/build_words.py [word-bank/letter-burn-word-bank.xlsx]

Reads the Answers sheet (Word, Display, Category, Accepted spellings, Check, Decision, Source)
and the Categories sheet. Rows marked Decision = drop are left out.
Needs openpyxl:  pip install openpyxl
"""
import json, os, re, sys, unicodedata
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "word-bank", "letter-burn-word-bank.xlsx")

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = s.lower().replace("'", "").replace(".", "").replace("-", " ")
    return re.sub(r"\s+", " ", s).strip()

wb = load_workbook(src, read_only=True, data_only=True)
cats = []
for r in wb["Categories"].iter_rows(min_row=2, values_only=True):
    if not r[0] or not r[1]:
        continue
    cats.append({"id": str(r[0]).strip(), "prompt": r[1], "group": r[2], "note": r[3], "minAnswers": int(r[5] or 0)})
cid = [c["id"] for c in cats]

words = {}
everyday = {}  # word -> categories from the hand-written lists (used for par, not loaded by the game)
for r in wb["Answers"].iter_rows(min_row=2, values_only=True):
    word, display, cat, spell, check, decision, source = (list(r) + [None] * 7)[:7]
    if not word or not cat or str(decision or "keep").strip().lower() == "drop":
        continue
    k = norm(word)
    if not re.fullmatch(r"[a-z ]+", k) or cat not in cid:
        print("skipped:", word, cat); continue
    w = words.setdefault(k, [display or word.capitalize(), [], set(), []])
    t = cid.index(cat)
    if t not in w[1]:
        w[1].append(t)
    if str(source or "list").strip().lower() == "list":
        everyday.setdefault(k, set()).add(t)
    for v in str(spell or "").split(","):
        v = norm(v)
        if v and v != k:
            w[2].add(v)
    if str(check or "").strip().lower() == "check" and t not in w[3]:
        w[3].append(t)

# compact format for the game: { word: [ [category numbers], [other spellings] ] }
out = {}
for k in sorted(words):
    d, tags, vs, chk = words[k]
    out[k] = [sorted(tags), sorted(vs)] if vs else [sorted(tags)]
for c in cats:
    c["answers"] = sum(1 for v in out.values() if cid.index(c["id"]) in v[0])

os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
json.dump({"categories": cid, "words": out}, open(os.path.join(ROOT, "data", "words.json"), "w"), separators=(",", ":"))
json.dump(cats, open(os.path.join(ROOT, "data", "categories.json"), "w"), indent=1)
json.dump({"categories": cid, "words": {k: sorted(v) for k, v in sorted(everyday.items())}},
          open(os.path.join(ROOT, "tools", "everyday.json"), "w"), separators=(",", ":"))
low = [f"{c['id']} ({c['answers']}/{c['minAnswers']})" for c in cats if c["answers"] < c["minAnswers"]]
print(f"{len(out)} words, {len(cats)} categories")
print("below minimum:", ", ".join(low) if low else "none")
